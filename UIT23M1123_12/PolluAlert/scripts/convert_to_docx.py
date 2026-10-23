import os
import re
import shutil
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "docs")
FIGURES_DIR = os.path.join(DOCS_DIR, "figures")
MD_FILE = os.path.join(DOCS_DIR, "PolluAlert_Project_Report.md")
DOCX_FILE = os.path.join(DOCS_DIR, "PolluAlert_Project_Report.docx")

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="94A3B8", sz="6"):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'  <w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def add_header_footer(doc):
    for s in doc.sections:
        # Header
        header = s.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run("PolluAlert: An NLP-Based Pollution Emergency Detection and Alert System")
        hr.font.name = "Calibri"
        hr.font.size = Pt(8.5)
        hr.font.color.rgb = RGBColor(148, 163, 184)
        
        # Footer
        footer = s.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run("Department of Computer Engineering — Academic Project Report")
        fr.font.name = "Calibri"
        fr.font.size = Pt(8.5)
        fr.font.color.rgb = RGBColor(148, 163, 184)

def add_styled_paragraph(doc, text, style='Normal', space_after=6, space_before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph(style=style)
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = 1.15
    
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**') and len(token) >= 4:
            run = p.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
            run = p.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(180, 40, 40)
        else:
            p.add_run(token)
    return p

def create_docx():
    print("Reading markdown report...")
    with open(MD_FILE, "r", encoding="utf-8") as f:
        lines = f.readlines()

    doc = Document()

    # Page setup - 1 inch standard margins
    for s in doc.sections:
        s.top_margin = Inches(1)
        s.bottom_margin = Inches(1)
        s.left_margin = Inches(1)
        s.right_margin = Inches(1)

    # Style defaults
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(30, 41, 59)

    add_header_footer(doc)

    in_code_block = False
    code_lines = []
    in_table = False
    table_lines = []

    def flush_table():
        nonlocal table_lines
        if not table_lines:
            return
        
        parsed_rows = []
        for tl in table_lines:
            cells = [c.strip() for c in tl.strip().strip('|').split('|')]
            if all(re.match(r'^:?-+:?$', c) for c in cells if c):
                continue
            parsed_rows.append(cells)
        
        if parsed_rows:
            col_count = max(len(r) for r in parsed_rows)
            table = doc.add_table(rows=len(parsed_rows), cols=col_count)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(table, color="94A3B8", sz="6")

            # Set header row to repeat on every page
            header_tr = table.rows[0]._element.get_or_add_trPr()
            header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

            for row_idx, row_data in enumerate(parsed_rows):
                is_header = (row_idx == 0)
                # Keep row together
                trPr = table.rows[row_idx]._element.get_or_add_trPr()
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

                for col_idx in range(col_count):
                    val = row_data[col_idx] if col_idx < len(row_data) else ""
                    cell = table.cell(row_idx, col_idx)
                    cell.text = ""
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                    cp = cell.paragraphs[0]
                    cp.paragraph_format.space_after = Pt(4)
                    cp.paragraph_format.space_before = Pt(4)
                    cp.paragraph_format.line_spacing = 1.05
                    cp.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    
                    tokens = re.split(r'(\*\*.*?\*\*)', val)
                    for token in tokens:
                        if not token:
                            continue
                        if token.startswith('**') and token.endswith('**'):
                            r = cp.add_run(token[2:-2])
                            r.bold = True
                        else:
                            r = cp.add_run(token)
                            if is_header:
                                r.bold = True

                    if is_header:
                        set_cell_background(cell, "1E3A8A")  # Deep Navy Blue
                        for r in cp.runs:
                            r.font.color.rgb = RGBColor(255, 255, 255)
                            r.font.size = Pt(10)
                            r.font.bold = True
                    else:
                        if row_idx % 2 == 1:
                            set_cell_background(cell, "F8FAFC")  # Subtle zebra striping
                        else:
                            set_cell_background(cell, "FFFFFF")
                        for r in cp.runs:
                            r.font.size = Pt(9.5)
                    set_cell_margins(cell, top=100, bottom=100, left=140, right=140)
            
            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(6)
        table_lines = []

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        code_text = "".join(code_lines)
        table = doc.add_table(rows=1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="CBD5E1", sz="4")
        cell = table.cell(0, 0)
        set_cell_background(cell, "F8FAFC")
        set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
        cp = cell.paragraphs[0]
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.space_before = Pt(2)
        run = cp.add_run(code_text.rstrip())
        run.font.name = 'Consolas'
        run.font.size = Pt(8.8)
        run.font.color.rgb = RGBColor(30, 41, 59)
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)
        code_lines = []

    idx = 0
    total = len(lines)
    print(f"Processing {total} lines with graphic embedding and structured tables...")

    while idx < total:
        line = lines[idx]
        stripped = line.strip()

        # Check code blocks
        if stripped.startswith("```"):
            if in_code_block:
                in_code_block = False
                flush_code()
            else:
                if in_table:
                    in_table = False
                    flush_table()
                in_code_block = True
            idx += 1
            continue

        if in_code_block:
            code_lines.append(line)
            idx += 1
            continue

        # Check table
        if stripped.startswith("|") and stripped.endswith("|"):
            if not in_table:
                in_table = True
            table_lines.append(stripped)
            idx += 1
            continue
        elif in_table:
            in_table = False
            flush_table()

        # Page breaks
        if stripped in ['<div style="page-break-after: always;"></div>', '\\newpage', '---']:
            if stripped == '---':
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
            else:
                doc.add_page_break()
            idx += 1
            continue

        if not stripped:
            idx += 1
            continue

        # Headings
        if stripped.startswith("# "):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(22)
            h.paragraph_format.space_after = Pt(10)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 58, 138)  # Deep Navy Blue
            idx += 1
            continue
        elif stripped.startswith("## "):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(8)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(2, 132, 199)  # Ocean Blue
            idx += 1
            continue
        elif stripped.startswith("### "):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(12)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59)
            idx += 1
            continue
        elif stripped.startswith("#### "):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(stripped[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(71, 85, 105)
            idx += 1
            continue

        # Check Images: ![alt](path)
        img_match = re.search(r'!\[(.*?)\]\((.*?)\)', stripped)
        if img_match:
            alt_text = img_match.group(1)
            img_path = img_match.group(2)
            
            # Resolve relative path
            full_img_path = os.path.normpath(os.path.join(DOCS_DIR, img_path))
            if os.path.exists(full_img_path):
                try:
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(10)
                    p.paragraph_format.space_after = Pt(4)
                    run = p.add_run()
                    # Scale width
                    run.add_picture(full_img_path, width=Inches(5.8))
                    
                    if alt_text:
                        cap = doc.add_paragraph()
                        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        cap.paragraph_format.space_before = Pt(2)
                        cap.paragraph_format.space_after = Pt(10)
                        cap.paragraph_format.keep_with_next = True
                        r = cap.add_run(f"Figure: {alt_text}")
                        r.italic = True
                        r.font.size = Pt(9.5)
                        r.font.color.rgb = RGBColor(71, 85, 105)
                except Exception as e:
                    print(f"Error adding image {full_img_path}: {e}")
            idx += 1
            continue

        # Check list items
        if stripped.startswith("- ") or stripped.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            content = stripped[2:]
            tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', content)
            for token in tokens:
                if not token:
                    continue
                if token.startswith('**') and token.endswith('**') and len(token) >= 4:
                    r = p.add_run(token[2:-2])
                    r.bold = True
                elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
                    r = p.add_run(token[1:-1])
                    r.italic = True
                elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
                    r = p.add_run(token[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                else:
                    p.add_run(token)
            idx += 1
            continue

        # Check numbered list
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            content = num_match.group(2)
            tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', content)
            for token in tokens:
                if not token:
                    continue
                if token.startswith('**') and token.endswith('**') and len(token) >= 4:
                    r = p.add_run(token[2:-2])
                    r.bold = True
                elif token.startswith('*') and token.endswith('*') and len(token) >= 2:
                    r = p.add_run(token[1:-1])
                    r.italic = True
                elif token.startswith('`') and token.endswith('`') and len(token) >= 2:
                    r = p.add_run(token[1:-1])
                    r.font.name = 'Consolas'
                    r.font.size = Pt(9.5)
                else:
                    p.add_run(token)
            idx += 1
            continue

        # Regular paragraph
        add_styled_paragraph(doc, stripped)
        idx += 1

    if in_table:
        flush_table()
    if in_code_block:
        flush_code()

    print(f"Saving compiled professional Word document to {DOCX_FILE}...")
    doc.save(DOCX_FILE)
    print("Document successfully compiled and saved!")

if __name__ == "__main__":
    create_docx()
